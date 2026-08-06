# -*- coding: utf-8 -*-
"""Shared library: fetch + parse + decode the survey Google Sheet.
Deterministic. No model tokens spent re-deriving this each wave.

Sheet layout (see project README):
  Employees: ID | ФИО | Отдел | Должность | Токен | Использование | Дата | Процент правильных |
             Балл самооценки | Средний балл | Принятие | Интерес | Безопасность | Портрет | Таймер
  Questions: ID вопроса | Вопрос | Включен | Ограничение (по ID пользователя)
  Answers:   ID вопроса | Вариант A..D | Номер правильного | Тег A..D
  Results:   ID сотрудника | ID вопроса | Ответ | Дата   (Ответ обычно буква A/B/C/D или свой текст)

Персональные наборы вопросов
----------------------------
Колонка «Ограничение (по ID пользователя)» на листе Questions задаёт, кому вопрос
показывается. Пусто = всем. Синтаксис разбирает `parse_restriction`, он ОБЯЗАН
совпадать с одноимённой функцией в Code.gs (бэкенд опроса) — иначе статистика
посчитает не тот набор, который человек реально видел.
"""
import os, re, json, urllib.request, ssl

HERE = os.path.dirname(os.path.abspath(__file__))
CFG = json.load(open(os.path.join(HERE, "config.json"), encoding="utf-8"))

RESTRICTION_COL = CFG.get("restriction_column", "Ограничение (по ID пользователя)")

# Маркер «не успел ответить» — фронтенд пишет его в Results, когда истёк таймер
# знаниевого блока. Пустая ячейка означает другое: вопрос пропущен молча.
TIMEOUT_MARKS = {"timeout", "не успел", "не успел ответить", "—"}

# Волны: номер пишется в эту колонку Results, архив итогов — на этом листе.
# Имена обязаны совпадать с WAVE_COLUMN / HISTORY_SHEET в Code.gs.
WAVE_COL = "Волна"
HISTORY_SHEET = "Employees_history"
_ALL_WORDS = {"все", "всем", "всё", "all", "*"}
_NONE_WORDS = {"никому", "никто", "нет", "none", "-"}
_EXCEPT_WORDS = ("кроме", "except", "!", "^")


def fetch_sheet(dest, sheet_id=None):
    """Download the live sheet as xlsx. Returns dest path."""
    sid = sheet_id or CFG["sheet_id"]
    url = CFG["export_url"].format(sheet_id=sid)
    ctx = ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = ssl.CERT_NONE
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, context=ctx, timeout=60) as r:
        data = r.read()
    with open(dest, "wb") as f:
        f.write(data)
    return dest


def _rows(ws):
    out = []
    for row in ws.iter_rows(values_only=True):
        vals = [("" if v is None else v) for v in row]
        out.append(vals)
    return out


def _qid(value):
    """'301.0' / 301 / ' 301 ' -> '301'."""
    s = str(value).strip()
    if not s:
        return ""
    return str(int(float(s))) if re.fullmatch(r"\d+(?:\.0+)?", s) else s


def _truthy(value):
    return value is True or str(value).strip().lower() in ("true", "да", "1", "yes")


# --------------------------------------------------------------------------
# Ограничение вопроса по ID сотрудника
# --------------------------------------------------------------------------
def parse_restriction(raw):
    """Разобрать ячейку «Ограничение (по ID пользователя)».

    Поддерживается (регистр не важен, разделители — запятая/точка с запятой/перевод строки):
      ''            — вопрос всем (по умолчанию)
      '2,5,7'       — только этим ID
      '2-8'         — диапазон ID
      '!5' '-5'     — всем, кроме 5
      'кроме 5,7'   — всем, кроме 5 и 7
      '2-8, !5'     — 2,3,4,6,7,8
      'все'         — всем явно;  'никому' — никому (вопрос выключен де-факто)

    Возвращает {'include': set|None, 'exclude': set, 'none': bool, 'raw': str, 'bad': [токены]}.
    include=None означает «все, кроме exclude».
    """
    rule = {"include": None, "exclude": set(), "none": False, "raw": str(raw or "").strip(), "bad": []}
    text = rule["raw"]
    if not text:
        return rule
    include = set()
    has_include = False
    default_negate = False   # «кроме» действует до конца выражения: «кроме 5, 7» = минус оба
    for token in re.split(r"[,;\n]+", text):
        t = token.strip().lower()
        if not t:
            continue
        if t in _ALL_WORDS:
            has_include = False
            include = set()
            continue
        if t in _NONE_WORDS:
            rule["none"] = True
            continue
        negate = default_negate
        for word in _EXCEPT_WORDS:
            if t.startswith(word):
                negate = True
                if word in ("кроме", "except"):
                    default_negate = True
                t = t[len(word):].strip()
                break
        if not negate and re.fullmatch(r"[-–]\s*\d+", t):     # «-5» = кроме 5
            negate = True
            t = t.lstrip("-–").strip()
        ids = _token_ids(t)
        if ids is None:
            rule["bad"].append(token.strip())
            continue
        if negate:
            rule["exclude"] |= ids
        else:
            has_include = True
            include |= ids
    if has_include:
        rule["include"] = include
    return rule


def _token_ids(t):
    """'5' -> {5};  '2-8' -> {2..8};  иначе None."""
    m = re.fullmatch(r"(\d+)\s*[-–—]\s*(\d+)", t)
    if m:
        lo, hi = int(m.group(1)), int(m.group(2))
        if lo > hi:
            lo, hi = hi, lo
        if hi - lo > 10000:
            return None
        return set(range(lo, hi + 1))
    if re.fullmatch(r"\d+", t):
        return {int(t)}
    return None


def restriction_allows(rule, emp_id):
    """Показывается ли вопрос с таким правилом сотруднику emp_id."""
    if rule.get("none"):
        return False
    try:
        eid = int(emp_id)
    except (TypeError, ValueError):
        return rule.get("include") is None
    if eid in rule.get("exclude", set()):
        return False
    include = rule.get("include")
    return True if include is None else eid in include


def restriction_label(rule):
    """Человекочитаемое описание правила — для отчёта."""
    if rule.get("none"):
        return "никому"
    inc, exc = rule.get("include"), sorted(rule.get("exclude", set()))
    if inc is None and not exc:
        return "всем"
    parts = []
    if inc is not None:
        parts.append("только ID " + _compact(sorted(inc)))
    if exc:
        parts.append(("кроме ID " if inc is None else "минус ID ") + _compact(exc))
    return ", ".join(parts) if parts else "всем"


def _compact(ids):
    """[2,3,4,7] -> '2-4,7'."""
    out, start, prev = [], None, None
    for i in list(ids) + [None]:
        if start is None:
            start = prev = i
            continue
        if i is not None and i == prev + 1:
            prev = i
            continue
        out.append(str(start) if start == prev else f"{start}-{prev}")
        start = prev = i
    return ",".join(out)


# --------------------------------------------------------------------------
def parse(xlsx_path):
    """Return {'employees', 'questions', 'question_order', 'answers', 'responses'}.

    questions: {qid: {'id','text','enabled','restriction','rule','correct'}} —
    порядок прохождения хранится отдельно в question_order.
    """
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    S = CFG["sheets"]

    # Employees
    erows = _rows(wb[S["employees"]])
    hdr = erows[0]
    ncol = max(i for i, h in enumerate(hdr) if str(h).strip()) + 1
    keys = [str(h).strip() for h in hdr[:ncol]]
    employees = []
    for r in erows[1:]:
        name = str(r[1]).strip() if len(r) > 1 else ""
        if not name:
            continue
        rec = {keys[i]: r[i] for i in range(min(ncol, len(r)))}
        employees.append(rec)

    # Answers -> options + correct letter
    answers = {}
    for r in _rows(wb[S["answers"]])[1:]:
        qid = _qid(r[0])
        if not qid:
            continue
        answers[qid] = {
            "A": r[1], "B": r[2], "C": r[3], "D": r[4],
            "correct": (str(r[5]).strip() or None) if len(r) > 5 else None,
            "tags": {"A": r[6] if len(r) > 6 else None, "B": r[7] if len(r) > 7 else None,
                     "C": r[8] if len(r) > 8 else None, "D": r[9] if len(r) > 9 else None},
        }

    # Questions -> текст, признак включённости, ограничение по ID
    questions, order = {}, []
    qsheet = wb[S["questions"]] if S.get("questions") in wb.sheetnames else None
    if qsheet is not None:
        qrows = _rows(qsheet)
        qh = [str(h).strip() for h in qrows[0]] if qrows else []
        def col(*names):
            for n in names:
                if n in qh:
                    return qh.index(n)
            return -1
        i_id, i_text = col("ID вопроса"), col("Вопрос")
        i_on = col("Включен", "Включено", "Включить")
        i_lim = col(RESTRICTION_COL, "Ограничение", "Ограничение (по ID пользователя)")
        for r in qrows[1:]:
            qid = _qid(r[i_id]) if i_id >= 0 and i_id < len(r) else ""
            if not qid:
                continue
            raw_lim = r[i_lim] if 0 <= i_lim < len(r) else ""
            questions[qid] = {
                "id": qid,
                "text": str(r[i_text]).strip() if 0 <= i_text < len(r) else "",
                "enabled": _truthy(r[i_on]) if 0 <= i_on < len(r) else True,
                "restriction": str(raw_lim or "").strip(),
                "rule": parse_restriction(raw_lim),
                "correct": (answers.get(qid) or {}).get("correct"),
            }
            order.append(qid)
    else:
        # Листа Questions нет — работаем по Answers, все вопросы всем.
        for qid in answers:
            questions[qid] = {"id": qid, "text": "", "enabled": True, "restriction": "",
                              "rule": parse_restriction(""), "correct": answers[qid].get("correct")}
            order.append(qid)

    # Results -> ответы по волнам: {wave: {employee_id: {qid: answer}}}
    # Колонка «Волна» может отсутствовать (таблица до внедрения волн) — тогда всё
    # считается волной 1, как и на стороне Code.gs.
    rs = wb[S["results"]]
    rhdr = [str(c).strip() if c is not None else "" for c in next(rs.iter_rows(max_row=1, values_only=True), ())]
    i_wave = rhdr.index(WAVE_COL) if WAVE_COL in rhdr else -1
    by_wave = {}
    for r in rs.iter_rows(min_row=2, values_only=True):
        if r[0] is None or r[1] is None:
            continue
        try:
            emp = str(int(float(r[0])))
        except (TypeError, ValueError):
            continue
        q = _qid(r[1])
        if not q:
            continue
        wave = 1
        if 0 <= i_wave < len(r) and r[i_wave] not in (None, ""):
            try:
                wave = int(float(r[i_wave]))
            except (TypeError, ValueError):
                wave = 1
        by_wave.setdefault(wave, {}).setdefault(emp, {})[q] = "" if r[2] is None else str(r[2]).strip()

    settings_map = _parse_settings(wb)
    current = _current_wave(settings_map, by_wave)
    return {"employees": employees, "questions": questions, "question_order": order,
            "answers": answers,
            "responses": by_wave.get(current, {}),   # ответы выбранной волны
            "responses_by_wave": by_wave, "wave": current, "waves": sorted(by_wave),
            "has_wave_column": i_wave >= 0,
            "settings": settings_map, "usages": _parse_usages(wb),
            "history": _parse_history(wb)}


def _current_wave(settings_map, by_wave):
    raw = settings_map.get("wave", "")
    try:
        n = int(float(str(raw).strip()))
        if n > 0:
            return n
    except (TypeError, ValueError):
        pass
    return max(by_wave) if by_wave else 1


def _parse_history(wb):
    """Лист Employees_history (пишет startNewWave) -> {wave: {emp_id: {колонка: значение}}}.

    Без него итоги прошлых волн в таблице не сохраняются — только ответы.
    """
    out = {}
    if HISTORY_SHEET not in wb.sheetnames:
        return out
    rows = _rows(wb[HISTORY_SHEET])
    if not rows:
        return out
    hdr = [str(h).strip() for h in rows[0]]
    if "Волна" not in hdr or "ID" not in hdr:
        return out
    i_w, i_id = hdr.index("Волна"), hdr.index("ID")
    for r in rows[1:]:
        if i_w >= len(r) or i_id >= len(r):
            continue
        try:
            wave, eid = int(float(r[i_w])), int(float(r[i_id]))
        except (TypeError, ValueError):
            continue
        out.setdefault(wave, {})[eid] = {hdr[i]: (r[i] if i < len(r) else "")
                                         for i in range(len(hdr)) if hdr[i]}
    return out


def select_wave(data, wave):
    """Переключить разбор на другую волну. Итоги берутся из Employees_history,
    если для этой волны есть архив, иначе из текущего Employees (и это надо назвать)."""
    data = dict(data)
    data["wave"] = wave
    data["responses"] = data.get("responses_by_wave", {}).get(wave, {})
    hist = data.get("history", {}).get(wave)
    data["scores_from_history"] = bool(hist)
    if hist:
        merged = []
        for e in data["employees"]:
            eid = emp_id(e)
            rec = dict(e)
            arch = hist.get(eid)
            # В архив могли скопировать всех подряд, включая тех, кто волну не проходил.
            # Прошедшим считаем только того, у кого в архиве есть дата прохождения:
            # иначе пустые строки раздуют число участников и испортят все доли.
            done = bool(arch) and str(arch.get("Дата и время прохождения", "")).strip() not in ("", "None")
            if arch:
                for k, v in arch.items():
                    if k not in ("Волна", "Архивировано") and v not in ("", None):
                        rec[k] = v
            rec["Использование"] = "Использован" if done else "Не использован"
            merged.append(rec)
        data["employees"] = merged
    return data


def _parse_settings(wb):
    """Лист Settings: A=ключ, B=значение. Значения-JSON разбираются автоматически."""
    out = {}
    if "Settings" not in wb.sheetnames:
        return out
    for r in wb["Settings"].iter_rows(values_only=True):
        if not r or r[0] is None:
            continue
        key = str(r[0]).strip()
        if not key:
            continue
        val = "" if len(r) < 2 or r[1] is None else str(r[1]).strip()
        if val[:1] in ("{", "["):
            try:
                val = json.loads(val)
            except ValueError:
                pass
        out[key] = val
    return out


def _parse_usages(wb):
    """Лист Usages -> {employee_id: {признак: 'Да'/'Нет'/'Не определено', ...}}.

    Связка по колонке ID (не по ФИО: тёзки и смена фамилии тихо ломают связь).
    Строки без ID пропускаются — их видно в analyze.py как усечённое покрытие.
    """
    out = {}
    if "Usages" not in wb.sheetnames:
        return out
    rows = _rows(wb["Usages"])
    if not rows:
        return out
    hdr = [str(h).strip() for h in rows[0]]
    if "ID" not in hdr:
        return out
    i_id = hdr.index("ID")
    for r in rows[1:]:
        if i_id >= len(r):
            continue
        raw = str(r[i_id]).strip()
        if not raw:
            continue
        try:
            eid = int(float(raw))
        except ValueError:
            continue
        out[eid] = {hdr[i]: ("" if i >= len(r) or r[i] is None else str(r[i]).strip())
                    for i in range(len(hdr)) if hdr[i]}
    return out


def settings(data, key, default=None):
    return data.get("settings", {}).get(key, default)


def min_answers(data):
    """Минимум ответов на характеристику. Единый источник — лист Settings (ключ
    `min_answers`, значение-JSON), иначе config.json. Должен совпадать с
    js/config.js -> portrait.dimensions.*.minAnswers."""
    fallback = CFG.get("portrait_min_answers", {"adoption": 6, "interest": 6, "safety": 4})
    val = settings(data, "min_answers")
    if isinstance(val, dict):
        merged = dict(fallback)
        for k, v in val.items():
            try:
                merged[k] = int(v)
            except (TypeError, ValueError):
                pass
        return merged
    return dict(fallback)


# --------------------------------------------------------------------------
# Вторая таблица: метрики, антиметрики, мероприятия по отделам
# --------------------------------------------------------------------------
def fetch_metrics(dest, sheet_id=None):
    """Скачать таблицу метрик. Читаем анонимно; писать в неё скилл не умеет."""
    return fetch_sheet(dest, sheet_id or CFG.get("metrics_sheet_id"))


def _norm_date(v):
    """Даты в шапке лежат и текстом ('01.08.2026'), и датой — приводим к ДД.ММ.ГГГГ."""
    if v in (None, ""):
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%d.%m.%Y")
    s = str(v).strip()
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})(?:[ T].*)?", s)
    return f"{m.group(3)}.{m.group(2)}.{m.group(1)}" if m else s


def parse_metrics(xlsx_path, dept_sheet=None):
    """Лист отдела -> {'blocks', 'antimetrics', 'events', 'dept_size'}.

    blocks: {'Метрики. Опрос': {'dates': [...], 'target_label': ..., 'rows': [...]}, ...}
    Каждая строка: {'name', 'source', 'values': {дата: значение}, 'target'}.
    """
    import openpyxl
    name = dept_sheet or CFG.get("metrics_dept_sheet", "БА")
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if name not in wb.sheetnames:
        return {"sheet": None, "available": wb.sheetnames, "blocks": {}, "antimetrics": [], "events": []}
    rows = _rows(wb[name])

    blocks, antimetrics, events, dept_size = {}, [], [], None
    current, i_target = None, None
    for r in rows:
        head = str(r[0]).strip() if r and r[0] not in (None, "") else ""

        if head.startswith("Метрики."):
            dates, i_target = [], None
            for i in range(2, len(r)):
                cell = str(r[i]).strip() if r[i] not in (None, "") else ""
                if cell.startswith("Цель"):
                    i_target = i
                    break
                if cell:
                    dates.append((i, _norm_date(r[i])))
            current = head
            blocks[current] = {"dates": [d for _, d in dates], "date_cols": [i for i, _ in dates],
                               "target_label": str(r[i_target]).strip() if i_target is not None else "",
                               "target_col": i_target, "rows": []}
        elif head == "Антиметрики":
            current = "Антиметрики"
        elif head and current and current.startswith("Метрики."):
            b = blocks[current]
            vals = {}
            for col, d in zip(b["date_cols"], b["dates"]):
                if col < len(r) and r[col] not in (None, ""):
                    vals[d] = r[col]
            b["rows"].append({
                "name": " ".join(head.split()),
                "source": str(r[1]).strip() if len(r) > 1 and r[1] not in (None, "") else "",
                "values": vals,
                "target": r[b["target_col"]] if b["target_col"] is not None and b["target_col"] < len(r) else None,
            })
        elif head and current == "Антиметрики":
            antimetrics.append({"name": " ".join(head.split()),
                                "source": str(r[1]).strip() if len(r) > 1 else "",
                                "trigger": str(r[2]).strip() if len(r) > 2 else "",
                                "action": str(r[3]).strip() if len(r) > 3 else ""})

        # Мероприятия и размер отдела живут в правых колонках того же листа.
        for i in range(len(r)):
            cell = str(r[i]).strip() if r[i] not in (None, "") else ""
            if cell == "Мероприятия" and i + 3 < len(r):
                events.append({"_header_col": i})
            elif cell == "Количество сотрудников в отделе" and i + 1 < len(r):
                dept_size = r[i + 1]

    col = events[0]["_header_col"] if events else None
    events = []
    if col is not None:
        started = False
        for r in rows:
            cell = str(r[col]).strip() if col < len(r) and r[col] not in (None, "") else ""
            if cell == "Мероприятия":
                started = True
                continue
            if started and cell:
                events.append({
                    "name": " ".join(cell.split()),
                    "status": str(r[col + 1]).strip() if col + 1 < len(r) and r[col + 1] not in (None, "") else "",
                    "date": _norm_date(r[col + 2]) if col + 2 < len(r) else "",
                    "result": str(r[col + 3]).strip() if col + 3 < len(r) and r[col + 3] not in (None, "") else "",
                })
    return {"sheet": name, "blocks": blocks, "antimetrics": antimetrics,
            "events": events, "dept_size": dept_size}


def emp_id(rec):
    try:
        return int(float(rec["ID"]))
    except (KeyError, TypeError, ValueError):
        return None


def completed_ids(data):
    """IDs of employees who finished (status contains 'Использован', not 'Не'), excluding test profiles."""
    test = set(CFG["test_employee_ids"])
    out = []
    for e in data["employees"]:
        eid = emp_id(e)
        status = str(e.get("Использование", "")).strip().lower()
        done = ("использован" in status) and ("не использован" not in status)
        if done and eid not in test:
            out.append(eid)
    return out


def not_completed(data):
    test = set(CFG["test_employee_ids"])
    out = []
    for e in data["employees"]:
        eid = emp_id(e)
        if eid in test:
            continue
        status = str(e.get("Использование", "")).strip().lower()
        if "не использован" in status or status == "":
            out.append({"id": eid, "name": e.get("ФИО"), "status": e.get("Использование")})
    return out


# --------------------------------------------------------------------------
# Наборы вопросов
# --------------------------------------------------------------------------
def questions_for(data, eid):
    """Список ID вопросов, которые ДОЛЖЕН был увидеть сотрудник eid (в порядке листа)."""
    return [q for q in data["question_order"]
            if data["questions"][q]["enabled"] and restriction_allows(data["questions"][q]["rule"], eid)]


def base_question_ids(data):
    """Все включённые вопросы без учёта ограничений — общий каркас волны."""
    return [q for q in data["question_order"] if data["questions"][q]["enabled"]]


def restricted_questions(data):
    """{qid: описание правила} только для вопросов с непустым ограничением."""
    out = {}
    for q in data["question_order"]:
        rec = data["questions"][q]
        if rec["restriction"]:
            out[q] = restriction_label(rec["rule"])
    return out


def knowledge_ids(data, asked=None):
    """Вопросы знаниевого блока (те, у кого в Answers указан правильный ответ)."""
    pool = asked if asked is not None else data["question_order"]
    return [q for q in pool if (data["questions"].get(q) or {}).get("correct")]


def knowledge_key(data):
    """{qid: правильная буква} из листа Answers; config.knowledge_correct — только запасной вариант."""
    key = {q: data["questions"][q]["correct"] for q in data["questions"] if data["questions"][q].get("correct")}
    return key or dict(CFG.get("knowledge_correct", {}))


def is_timeout(value):
    return str(value).strip().lower() in TIMEOUT_MARKS


def knowledge_score(resp, asked_ids, key):
    """{correct, total, wrong, blank, timeout} по НАБОРУ ЭТОГО человека.

    «Не успел» (истёк таймер) и «пропустил молча» считаются неверным ответом —
    так определена метрика проекта, — но разносятся по разным спискам:
    не знал ≠ не успел, и на воркшопе это разные выводы.
    """
    ids = [q for q in asked_ids if q in key]
    wrong, blank, timeout = [], [], []
    for q in ids:
        got = resp.get(q, "")
        if is_timeout(got):
            timeout.append(q)
            wrong.append(q)
        elif got == "":
            blank.append(q)
            wrong.append(q)
        elif got != key[q]:
            wrong.append(q)
    return {"correct": len(ids) - len(wrong), "total": len(ids),
            "wrong": wrong, "blank": blank, "timeout": timeout}


def decode_answer(qid, resp):
    """Human-readable meaning of a preference/attitude answer, or the free text."""
    val = resp.get(qid, "")
    if is_timeout(val):
        return "(не успел ответить)"
    d = CFG["decode"].get(qid)
    if val in ("A", "B", "C", "D") and d and val in d:
        return d[val]
    if qid == "511" and val in CFG["self_level"]:
        return CFG["self_level"][val]
    return "(свой ответ) " + val if val else "—"


def answer_state(qid, resp, asked_ids):
    """Различает «не задавали», «не успел» и «пропустил» — для честных знаменателей."""
    if qid not in asked_ids:
        return "не задавался"
    val = resp.get(qid, "")
    if is_timeout(val):
        return "не успел"
    if val == "":
        return "без ответа"
    return "ok"
